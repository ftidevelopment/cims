from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class BaseExportService:

    HEADER_FONT = Font(bold=True)

    HEADER_FILL = PatternFill(
        fill_type="solid",
        start_color="D9EAD3",
        end_color="D9EAD3",
    )

    HEADER_ALIGNMENT = Alignment(
        horizontal="center",
        vertical="center",
    )

    def __init__(self, repository):
        self.repository = repository

    def export(
        self,
        sheet_name,
        export_config,
        keyword=None,
        is_active=True,
        sort_by=None,
        sort_order="asc",
    ):
        rows = self.repository.get_export_data(
            keyword=keyword,
            is_active=is_active,
            sort_by=sort_by,
            sort_order=sort_order,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = sheet_name

        # ===========================
        # Header
        # ===========================
        for column_index, column in enumerate(export_config, start=1):

            cell = worksheet.cell(
                row=1,
                column=column_index,
            )

            cell.value = column["header"]
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT

        # ===========================
        # Data
        # ===========================
        for row_index, row in enumerate(rows, start=2):

            for column_index, column in enumerate(export_config, start=1):

                field = column["field"]

                value = getattr(
                    row,
                    field,
                    ""
                )

                # ======================================
                # Boolean Display
                # ======================================
                value = self.format_export_value(
                    field,
                    value,
                )

                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).value = value

        # ===========================
        # Column Width
        # ===========================
        for index, column in enumerate(export_config, start=1):

            worksheet.column_dimensions[
                get_column_letter(index)
            ].width = column.get("width", 20)

        # ===========================
        # Freeze Header
        # ===========================
        worksheet.freeze_panes = "A2"

        # ===========================
        # Auto Filter
        # ===========================
        worksheet.auto_filter.ref = worksheet.dimensions

        return workbook

    def format_export_value(
        self,
        field,
        value,
    ):
        # ======================================
        # Active / Inactive
        # ======================================
        if field == "is_active":
            return "Active" if bool(value) else "Inactive"

        # ======================================
        # Locked / Normal
        # ======================================
        if field == "is_locked":
            return "Locked" if bool(value) else "Normal"

        return value